from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock, patch
from openpyxl import Workbook

from slopekg.inspection_history import (
    SOURCE, build_inspection_history, history_catalog, history_interfaces, history_payload, history_source,
)
from slopekg.server import SlopeKGHandler
from tests.support import workspace_paths


class InspectionHistoryTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.paths = workspace_paths(Path(self.directory.name))
        self.file = self.paths.root / SOURCE
        self.file.parent.mkdir(parents=True)
        book = Workbook()
        slope = book.active
        slope.title = '边坡'
        slope.append(['编号', '时间', '起止桩号', '长度m', '方向及位置', '存在问题', None])
        slope.append([1, '2022-04-25 00:00:00', 'K340+540～K340+600', 60, '下行', '路段内边坡水毁', 'unlabelled'])
        slope.append([2, 'invalid', 'K340+540', 0, '上行', '<img src=x onerror=alert(1)>'])
        bridge = book.create_sheet('桥梁')
        bridge.append(['编号', '时间', '起止桩号', '桥名', '方向及位置', '存在问题'])
        bridge.append([1, '2022-03-22 00:00:00', 'K312+073', None, '下行', '桥头水毁'])
        book.create_sheet('Sheet3')
        book.save(self.file)
        book.close()
        self.catalog = build_inspection_history(self.paths)

    def test_keeps_original_cells_and_does_not_infer_a_slope_from_a_station_range(self):
        rows = history_payload({'asset_type': ['all']}, self.paths)['rows']
        self.assertEqual(self.catalog['counts_by_type'], {'边坡': 2, '桥梁': 1})
        slope = next(r for r in rows if r['station_text'].endswith('600'))
        self.assertIsNone(slope['slope_id'])
        self.assertIsNone(slope['route_id'])
        self.assertEqual(slope['raw_fields']['未命名列G'], 'unlabelled')
        self.assertEqual(slope['source']['cells']['存在问题'], 'F2')
        self.assertFalse(slope['risk_eligible'])
        self.assertEqual(self.catalog['invalid_date_rows'], 1)
        self.assertEqual(self.catalog['skipped_sheets'], ['Sheet3'])

    def test_separate_types_date_keyword_and_pagination(self):
        query = {'asset_type': ['all'], 'q': ['水毁'], 'start': ['2022-04-01'], 'end': ['2022-04-30']}
        rows = history_payload(query, self.paths)['rows']
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['asset_type'], '边坡')
        self.assertEqual(history_payload({'asset_type': ['桥梁']}, self.paths)['total'], 1)
        self.assertEqual(history_payload({'asset_type': ['all'], 'limit': ['1'], 'offset': ['1']}, self.paths)['count'], 1)
        for bad in [{'limit': ['101']}, {'offset': ['-1']}, {'asset_type': ['slope']}, {'start': ['bad']}, {'start': ['2023-01-01'], 'end': ['2022-01-01']}]:
            with self.subTest(bad=bad), self.assertRaises(ValueError):
                history_payload(bad, self.paths)

    def test_same_source_reimport_has_stable_ids_and_no_duplicate_records(self):
        before = history_payload({'asset_type': ['all']}, self.paths)['rows']
        again = build_inspection_history(self.paths)
        self.assertEqual(again['batch_id'], self.catalog['batch_id'])
        self.assertEqual(before, history_payload({'asset_type': ['all']}, self.paths)['rows'])
        self.assertEqual(again['count'], 3)

    def test_undated_records_are_explicitly_included_without_inventing_a_date(self):
        query = {'start': ['2022-04-01'], 'end': ['2022-04-30']}
        self.assertEqual(history_payload(query, self.paths)['total'], 1)
        result = history_payload({**query, 'include_undated': ['true']}, self.paths)
        self.assertEqual(result['total'], 2)
        self.assertIsNone(result['rows'][0]['observed_date'])
        with self.assertRaises(ValueError):
            history_payload({'include_undated': ['yes']}, self.paths)

    def test_changed_source_is_not_downloaded_as_the_old_version(self):
        _, data = history_source(self.paths)
        self.assertEqual(data, self.file.read_bytes())
        self.file.write_bytes(b'changed source')
        with self.assertRaises(OSError):
            history_source(self.paths)

    def test_changed_generated_history_is_rejected(self):
        file = self.paths.inspection_history_dir / self.catalog['events_file']
        file.write_text('{"rows": []}', encoding='utf-8')
        with self.assertRaises(OSError):
            history_payload({}, self.paths)

    def test_bad_import_keeps_published_history_and_recent_interface_reserved(self):
        self.file.write_bytes(b'not an Excel file')
        with self.assertRaises(Exception):
            build_inspection_history(self.paths)
        self.assertEqual(history_catalog(self.paths)['count'], 3)
        original = {'interfaces': [{'path': '/api/inspections', 'status': 'reserved_not_connected'}]}
        result = history_interfaces(original, self.paths)
        self.assertEqual(result['interfaces'][0], original['interfaces'][0])
        self.assertEqual(next(r for r in result['interfaces'] if r['path'] == '/api/inspections/history')['status'], 'implemented_historical')

    def test_http_history_route_and_query_errors(self):
        with patch('slopekg.server.PATHS', self.paths):
            for url, status in [('/api/inspections/history?asset_type=桥梁', None), ('/api/inspections/history?limit=0', 400)]:
                handler = object.__new__(SlopeKGHandler)
                handler.path, handler.pipeline, handler.write_json = url, Mock(), Mock()
                handler.do_GET()
                if status:
                    self.assertEqual(handler.write_json.call_args.kwargs['status'], status)
                else:
                    self.assertEqual(handler.write_json.call_args.args[0]['total'], 1)

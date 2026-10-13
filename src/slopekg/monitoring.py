"""Source-preserving Huichang monitoring pilot; measurements are not risk labels."""
from __future__ import annotations

from collections import Counter, defaultdict
from contextlib import closing
from datetime import date, datetime, timedelta, timezone
import hashlib
import json
import math
from pathlib import Path
import sqlite3
import statistics
import tempfile
from io import BytesIO
from typing import Any

from .config import DemoPaths, PATHS
from .storage import read_json, write_json, verify_file_sha256


SCHEMA_VERSION = "huichang-monitoring-2"
PILOT_FILES = (
    "09YL01(雨量)单日.xlsx", "02GP01(地表位移)累计位移_日.xlsx",
    "02QJ02(倾角)累计角度_日.xlsx", "05LF03(裂缝)累计位移_日.xlsx",
    "10SW01(深部位移).xlsx", "07HS01(土壤含水率).xlsx",
    "07DX02(地下水位).xlsx", "02JS03(加速度).xlsx",
)
BOUNDARY = "历史原始监测回放；单位、分量、基准及测点归属待核；不输出风险等级或失稳预测。"


def _digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _timestamp(value: Any) -> str | None:
    try:
        parsed = value if isinstance(value, datetime) else datetime.fromisoformat(str(value).strip())
        if parsed.tzinfo is not None:
            raise ValueError("源时间含时区，需另行定义转换口径")
        return parsed.isoformat(sep=" ", timespec="seconds")
    except (TypeError, ValueError):
        return None


def _components(value: Any) -> list[float] | None:
    try:
        numbers = [float(part.strip()) for part in str(value).replace("，", ",").split(",")]
        return numbers if all(math.isfinite(x) for x in numbers) else None
    except (TypeError, ValueError):
        return None


def _raw(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    return value


def _blocks(days: set[str]) -> list[dict[str, Any]]:
    result = []
    for day in sorted(days):
        if result and (date.fromisoformat(day) - date.fromisoformat(result[-1]["end"])).days == 1:
            result[-1]["end"] = day
            result[-1]["days"] += 1
        else:
            result.append({"start": day, "end": day, "days": 1})
    return result


def _read_series(source: Path, root: Path) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    from openpyxl import load_workbook

    data = source.read_bytes()
    sha = _digest(data)
    relative = source.resolve().relative_to(root.resolve()).as_posix()
    series_id = "ms_" + _digest(f"{relative}:{sha}".encode())[:24]
    book = load_workbook(BytesIO(data), read_only=True, data_only=False)
    observations = []
    sheet_headers = []
    device_names = set()
    try:
        for sheet in book:
            iterator = sheet.iter_rows(values_only=True)
            headers = [str(x or "").strip() for x in next(iterator, ())]
            if "时间" not in headers:
                raise ValueError(f"{source.name}/{sheet.title}缺少时间字段")
            time_col = headers.index("时间")
            value_cols = [i for i, h in enumerate(headers) if h in {"数值", "单日雨量值", "累计雨量值"} or "(%)" in h]
            if not value_cols:
                raise ValueError(f"{source.name}/{sheet.title}缺少观测字段")
            sheet_headers.append({"sheet": sheet.title, "headers": headers})
            device_col = headers.index("设备名称") if "设备名称" in headers else None
            for row_number, row in enumerate(iterator, 2):
                if not any(x is not None for x in row):
                    continue
                stamp = _timestamp(row[time_col])
                raw_values = {headers[i]: _raw(row[i]) for i in value_cols}
                fields = {headers[i]: _components(row[i]) for i in value_cols}
                flags = []
                if stamp is None:
                    flags.append("invalid_time")
                if any(v is None for v in fields.values()):
                    flags.append("invalid_value")
                if any(v and any(n == 0 for n in v) for v in fields.values()):
                    flags.append("contains_zero")
                device = str(row[device_col]) if device_col is not None else source.stem.split("(")[0]
                device_names.add(device)
                observations.append({
                    "id": f"{series_id}:{sheet.title}:{row_number}", "series_id": series_id,
                    "time": stamp, "day": stamp[:10] if stamp else None,
                    "raw_time": _raw(row[time_col]), "device_name": device,
                    "raw_values": raw_values, "fields": fields,
                    "components": fields[headers[value_cols[0]]], "quality_flags": flags,
                    "source": {"file": relative, "sha256": sha, "sheet": sheet.title, "row": row_number,
                               "cells": {h: f"{_column_letter(i + 1)}{row_number}" for i, h in enumerate(headers) if i == time_col or i in value_cols}},
                })
    finally:
        book.close()
    if not observations:
        raise ValueError(f"{source.name}没有观测记录")
    time_counts = Counter(o["time"] for o in observations if o["time"])
    for observation in observations:
        if observation["time"] and time_counts[observation["time"]] > 1:
            observation["quality_flags"].append("duplicate_time")
    ordered = sorted((o for o in observations if o["time"]), key=lambda o: (o["time"], o["source"]["row"]))
    times = sorted(time_counts)
    gaps = [(datetime.fromisoformat(b) - datetime.fromisoformat(a)).total_seconds() for a, b in zip(times, times[1:])]
    days = {o["day"] for o in ordered}
    widths = Counter(len(o["components"]) for o in observations if o["components"] is not None)
    width = max(widths, default=0)
    stats = []
    for index in range(width):
        valid = [(o, o["components"][index]) for o in ordered if o["components"] is not None and len(o["components"]) > index]
        values = [v for _, v in valid]
        jumps = [(abs(b[1] - a[1]), a, b) for a, b in zip(valid, valid[1:]) if a[0]["time"] != b[0]["time"]]
        biggest = max(jumps, key=lambda j: j[0]) if jumps else None
        stats.append({"component": index + 1, "min": min(values) if values else None, "max": max(values) if values else None,
                      "zero_count": sum(v == 0 for v in values), "valid_count": len(values),
                      "largest_adjacent_change": {"change": biggest[0], "from": biggest[1][0]["time"], "to": biggest[2][0]["time"],
                                                  "from_value": biggest[1][1], "to_value": biggest[2][1],
                                                  "source_row": biggest[2][0]["source"]["row"]} if biggest else None})
    category = source.stem.split("(", 1)[1].split(")", 1)[0]
    unit = "%" if any("(%)" in h for s in sheet_headers for h in s["headers"]) else None
    meta = {
        "id": series_id, "sensor_id": "mp_" + _digest(f"{source.parent.name}:{source.stem.split('(')[0]}".encode())[:20],
        "site_id": "monitoring_site_huichang_202610", "slope_id": None, "file_name": source.name,
        "source_file": relative, "sha256": sha, "source_folder": source.parent.name, "headers": sheet_headers,
        "category": category, "device_names": sorted(device_names), "unit": unit,
        "unit_status": "source_header" if unit else "unknown", "component_names": [f"分量{i + 1}" for i in range(width)],
        "component_meaning_status": "unknown", "timezone": None, "baseline": None,
        "row_count": len(observations), "first_time": times[0] if times else None, "last_time": times[-1] if times else None,
        "recorded_days": len(days), "continuous_blocks": _blocks(days),
        "median_interval_seconds": statistics.median(gaps) if gaps else None,
        "max_interval_seconds": max(gaps) if gaps else None,
        "duplicate_time_rows": sum(n for n in time_counts.values() if n > 1),
        "invalid_time_rows": sum(o["time"] is None for o in observations),
        "invalid_value_rows": sum("invalid_value" in o["quality_flags"] for o in observations),
        "component_width_counts": dict(widths), "component_stats": stats,
        "review_status": "pending", "risk_eligible": False,
        "notes": ["按来源记录关联候选工点，逐设备对应待确认", "日记录日期不代表数据发布时间", "数值变化只作回查线索"],
    }
    return meta, observations


def _column_letter(index: int) -> str:
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result


def build_huichang_pilot(paths: DemoPaths = PATHS) -> dict[str, Any]:
    source_dir = paths.root / "data" / "202610" / "监测数据" / "会昌"
    sources = [source_dir / name for name in PILOT_FILES]
    missing = [p.name for p in sources if not p.is_file()]
    if missing:
        raise FileNotFoundError("会昌样例缺少文件：" + "、".join(missing))
    for source in sources:
        source.resolve().relative_to(source_dir.resolve())
    pairs = [_read_series(p, paths.root) for p in sources]
    batch_id = _digest((SCHEMA_VERSION + json.dumps([(m["source_file"], m["sha256"]) for m, _ in pairs])).encode())[:24]
    run_id = "monitoring_run_" + batch_id
    for meta, observations in pairs:
        meta["processing_run_id"] = run_id
        for observation in observations:
            observation["processing_run_id"] = run_id
    folder = paths.monitoring_dir / "snapshots" / batch_id
    folder.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=folder, suffix=".sqlite", delete=False) as handle:
        staging = Path(handle.name)
    try:
        with closing(sqlite3.connect(staging)) as connection, connection:
            connection.execute("CREATE TABLE observations (series_id TEXT, stamp TEXT, day TEXT, source_row INTEGER, payload TEXT)")
            connection.execute("CREATE INDEX observation_lookup ON observations(series_id, day, stamp)")
            for _, observations in pairs:
                connection.executemany("INSERT INTO observations VALUES (?, ?, ?, ?, ?)",
                                       ((o["series_id"], o["time"], o["day"], o["source"]["row"], json.dumps(o, ensure_ascii=False, allow_nan=False)) for o in observations))
        # A new immutable generation is published through one atomic catalog replacement.
        database_sha = _digest(staging.read_bytes())
        database = folder / f"observations-{database_sha[:16]}.sqlite"
        if database.exists():
            if _digest(database.read_bytes()) != database_sha:
                raise ValueError("已有监测数据库内容与其校验和不一致")
            staging.unlink()
        else:
            staging.replace(database)
    finally:
        staging.unlink(missing_ok=True)
    common = set.intersection(*({o["day"] for o in obs if o["day"]} for _, obs in pairs))
    catalog = {
        "status": "imported_pending_interpretation", "schema_version": SCHEMA_VERSION, "batch_id": batch_id,
        "generated_at": datetime.now(timezone.utc).isoformat(), "boundary": BOUNDARY,
        "processing_run": {"id": run_id, "software": "slopekg.monitoring", "version": SCHEMA_VERSION,
                           "parameters": {"source_files": list(PILOT_FILES), "imputation": "none", "unit_inference": False,
                                          "replay_aggregation": "last recorded observation per source date"}},
        "site": {"id": "monitoring_site_huichang_202610", "label": "会昌监测样例", "route_candidate": "G357 K399+175—K399+320",
                 "mapping_status": "candidate_from_document", "mapping_source": "data/202610/监测点布设.docx", "slope_id": None},
        "database": database.relative_to(paths.monitoring_dir).as_posix(), "database_sha256": database_sha,
        "files": len(pairs), "observations": sum(len(obs) for _, obs in pairs), "series": [m for m, _ in pairs],
        "common_days": len(common), "common_blocks": _blocks(common),
        "default_start": "2026-05-25", "default_end": "2026-09-01", "risk_eligible": False,
    }
    write_json(folder / "catalog.json", catalog)
    write_json(paths.monitoring_dir / "catalog.json", catalog)
    return catalog


def monitoring_catalog(paths: DemoPaths = PATHS) -> dict[str, Any]:
    catalog = read_json(paths.monitoring_dir / "catalog.json", None)
    if not catalog:
        return {"status": "not_imported", "series": [], "files": 0, "observations": 0, "boundary": BOUNDARY, "risk_eligible": False}
    database = (paths.monitoring_dir / catalog["database"]).resolve()
    database.relative_to(paths.monitoring_dir.resolve())
    if not database.is_file():
        raise FileNotFoundError("监测目录已生成但观测数据库缺失，请重新构建样例")
    verify_file_sha256(database, catalog.get("database_sha256"))
    return catalog


def public_catalog(paths: DemoPaths = PATHS) -> dict[str, Any]:
    return {k: v for k, v in monitoring_catalog(paths).items() if k not in {"database", "database_sha256"}}


def _dates(query: dict[str, list[str]], catalog: dict[str, Any]) -> tuple[str, str]:
    start = query.get("start", [catalog.get("default_start", "2026-05-25")])[0]
    end = query.get("end", [catalog.get("default_end", "2026-09-01")])[0]
    for value in (start, end):
        if date.fromisoformat(value).isoformat() != value:
            raise ValueError("日期须为YYYY-MM-DD")
    if start > end or (date.fromisoformat(end) - date.fromisoformat(start)).days > 2000:
        raise ValueError("日期范围须有序且不超过2001天")
    return start, end


def _connection(catalog: dict[str, Any], paths: DemoPaths) -> sqlite3.Connection:
    database = (paths.monitoring_dir / catalog["database"]).resolve()
    database.relative_to(paths.monitoring_dir.resolve())
    return sqlite3.connect(database.as_uri() + "?mode=ro", uri=True)


def observation_payload(query: dict[str, list[str]], paths: DemoPaths = PATHS) -> dict[str, Any]:
    catalog = monitoring_catalog(paths)
    if not catalog.get("series"):
        return {"status": "not_imported", "rows": [], "count": 0, "total": 0, "risk_eligible": False}
    series_id = query.get("series_id", [catalog["series"][0]["id"]])[0]
    if not any(s["id"] == series_id for s in catalog["series"]):
        raise KeyError("unknown monitoring series")
    start, end = _dates(query, catalog)
    limit = int(query.get("limit", ["100"])[0])
    offset = int(query.get("offset", ["0"])[0])
    if not 1 <= limit <= 1000 or offset < 0:
        raise ValueError("limit须为1—1000，offset须为非负整数")
    include_invalid = query.get("include_invalid_time", ["false"])[0]
    if include_invalid not in {"true", "false"}:
        raise ValueError("include_invalid_time须为true或false")
    where = "series_id=? AND (day>=? AND day<=?" + (" OR day IS NULL)" if include_invalid == "true" else ")")
    with closing(_connection(catalog, paths)) as connection:
        total = connection.execute(f"SELECT COUNT(*) FROM observations WHERE {where}", (series_id, start, end)).fetchone()[0]
        records = connection.execute(f"SELECT payload FROM observations WHERE {where} ORDER BY stamp,source_row LIMIT ? OFFSET ?", (series_id, start, end, limit, offset)).fetchall()
    return {"status": catalog["status"], "batch_id": catalog["batch_id"], "rows": [json.loads(r[0]) for r in records],
            "count": len(records), "total": total, "offset": offset, "limit": limit, "start": start, "end": end,
            "risk_eligible": False, "time_semantics": "按原始记录日期筛选，数据发布时间未知"}


def replay_payload(query: dict[str, list[str]], paths: DemoPaths = PATHS) -> dict[str, Any]:
    catalog = monitoring_catalog(paths)
    if not catalog.get("series"):
        return {"status": "not_imported", "series": [], "dates": [], "risk_eligible": False}
    start, end = _dates(query, catalog)
    dates = [(date.fromisoformat(start) + timedelta(days=i)).isoformat() for i in range((date.fromisoformat(end) - date.fromisoformat(start)).days + 1)]
    series = []
    with closing(_connection(catalog, paths)) as connection:
        for meta in catalog["series"]:
            records = connection.execute("SELECT payload FROM observations WHERE series_id=? AND day>=? AND day<=? ORDER BY stamp,source_row", (meta["id"], start, end)).fetchall()
            by_day: dict[str, list[dict]] = defaultdict(list)
            for record in records:
                obs = json.loads(record[0])
                by_day[obs["day"]].append(obs)
            points = []
            for day in dates:
                observations = by_day.get(day, [])
                last = observations[-1] if observations else None
                invalid = any("invalid_value" in o["quality_flags"] or "duplicate_time" in o["quality_flags"] for o in observations)
                points.append({"day": day, "count": len(observations), "components": last["components"] if last and not invalid else None,
                               "last_record": last, "quality_flags": ["ambiguous_or_invalid_day"] if invalid else ([] if last else ["no_record"])} )
            series.append({"meta": meta, "points": points, "record_count": len(records)})
    return {"status": catalog["status"], "batch_id": catalog["batch_id"], "start": start, "end": end, "dates": dates, "series": series,
            "aggregation": "每个原始记录日期的最后一条已记录观测；保留无记录日，不插值；不是官方日统计定义",
            "boundary": BOUNDARY, "risk_eligible": False}


def monitoring_graph(paths: DemoPaths = PATHS) -> dict[str, Any]:
    catalog = monitoring_catalog(paths)
    if not catalog.get("series"):
        return {"meta": {"status": "not_imported"}, "nodes": [], "edges": []}
    site = catalog["site"]
    run = catalog.get("processing_run", {"id": "monitoring_run_" + catalog["batch_id"], "version": catalog["schema_version"]})
    nodes = [{"id": site["id"], "type": "MonitoringSite", "label": site["label"], "props": site},
             {"id": run["id"], "type": "ProcessingRun", "label": "监测导入与质检版本", "props": run}]
    edges = []
    for meta in catalog["series"]:
        asset = "mf_" + meta["id"]
        quality = "mq_" + meta["id"]
        nodes.extend([
            {"id": meta["sensor_id"], "type": "MonitoringPoint", "label": "/".join(meta["device_names"]), "props": {"site_id": site["id"], "review_status": "pending"}},
            {"id": meta["id"], "type": "MonitoringSeries", "label": meta["category"], "props": meta},
            {"id": asset, "type": "DataTableAsset", "label": meta["file_name"], "props": {"file": meta["source_file"], "sha256": meta["sha256"]}},
            {"id": quality, "type": "QualityReport", "label": "数据核查", "props": {k: meta[k] for k in ("invalid_time_rows", "invalid_value_rows", "duplicate_time_rows", "risk_eligible")}},
        ])
        for source, target, relation in ((site["id"], meta["sensor_id"], "HAS_MONITORING_POINT"), (meta["sensor_id"], meta["id"], "HAS_MONITORING_SERIES"),
                                        (meta["id"], asset, "DERIVED_FROM"), (meta["id"], quality, "HAS_QUALITY_REPORT"),
                                        (meta["id"], run["id"], "GENERATED_BY"), (quality, run["id"], "GENERATED_BY")):
            edges.append({"source": source, "target": target, "type": relation, "review_status": "source_structure_only"})
    return {"meta": {"scope": "huichang_monitoring_pilot", "boundary": BOUNDARY, "batch_id": catalog["batch_id"]}, "nodes": nodes, "edges": edges}


def source_file(series_id: str, paths: DemoPaths = PATHS) -> Path:
    meta = next((s for s in monitoring_catalog(paths).get("series", []) if s["id"] == series_id), None)
    if meta is None:
        raise KeyError("unknown monitoring series")
    source = (paths.root / meta["source_file"]).resolve()
    source.relative_to((paths.root / "data" / "202610" / "监测数据" / "会昌").resolve())
    if _digest(source.read_bytes()) != meta["sha256"]:
        raise ValueError("原表已变化，与当前导入版本不一致，请重新构建样例")
    return source


def monitoring_interfaces(payload: dict[str, Any], paths: DemoPaths = PATHS) -> dict[str, Any]:
    try:
        catalog = monitoring_catalog(paths)
    except (OSError, ValueError, KeyError, TypeError) as exc:
        rows = [{**row, "status": "data_unavailable", "description": "会昌样例暂不可读：" + str(exc)}
                if row.get("path") == "/api/monitoring/observations" else dict(row)
                for row in payload.get("interfaces", [])]
        return {**payload, "interfaces": rows, "status_counts": dict(Counter(r["status"] for r in rows))}
    if not catalog.get("series"):
        return payload
    rows = [{**row, "status": "implemented_pilot", "scope": "会昌八表历史样例；字段待核", "description": "原始观测分页查询；按日期过滤，保留原表行号与未知字段"} if row.get("path") == "/api/monitoring/observations" else dict(row) for row in payload.get("interfaces", [])]
    return {**payload, "interfaces": rows, "status_counts": dict(Counter(r["status"] for r in rows))}


def export_monitoring_report(destination: Path, query: dict[str, list[str]], paths: DemoPaths = PATHS) -> dict[str, Any]:
    catalog = public_catalog(paths)
    replay = replay_payload(query, paths)
    if not replay.get("series"):
        raise ValueError("尚未导入会昌监测样例")
    payload = {"catalog": catalog, "replay": replay, "graph": monitoring_graph(paths)}
    encoded = json.dumps(payload, ensure_ascii=False, allow_nan=False).replace("<", "\\u003c").replace("\u2028", "\\u2028").replace("\u2029", "\\u2029")
    template = (paths.web_dir / "monitoring.html").read_text(encoding="utf-8")
    css = (paths.web_dir / "pages.css").read_text(encoding="utf-8") + "\n" + (paths.web_dir / "monitoring.css").read_text(encoding="utf-8")
    api = (paths.web_dir / "api.js").read_text(encoding="utf-8").replace("</script", "<\\/script")
    js = (paths.web_dir / "monitoring.js").read_text(encoding="utf-8").replace("</script", "<\\/script")
    template = template.replace('<link rel="stylesheet" href="./pages.css"/>', f"<style>{css}</style>")
    template = template.replace('<link rel="stylesheet" href="./monitoring.css"/>', "")
    template = template.replace('<script src="./api.js"></script><script src="./monitoring.js"></script>', f"<script>window.MONITORING_OFFLINE={encoded};</script><script>{api}</script><script>{js}</script>")
    destination.parent.mkdir(parents=True, exist_ok=True)
    from .storage import _atomic_write_text
    _atomic_write_text(destination, template)
    return {"status": "exported", "series": len(replay["series"]), "start": replay["start"], "end": replay["end"], "bytes": destination.stat().st_size, "risk_eligible": False}

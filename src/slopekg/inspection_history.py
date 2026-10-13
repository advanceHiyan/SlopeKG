"""Historical patrol evidence, separate from current slope/risk observations."""
from __future__ import annotations

from collections import Counter
from datetime import date, datetime, timezone
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
from typing import Any

from .config import PATHS, DemoPaths
from .storage import read_json, verify_file_sha256, write_json

VERSION = 'inspection-history-1'
SOURCE = 'data/202610/其他/日常巡查记录.xlsx'
BOUNDARY = '历史巡查记录；路线与具体边坡对应尚未提供。用于病害检索和原表核对，不作为近期状态或风险等级。'
HISTORY_INTERFACES = [
    {'path': '/api/inspections/history', 'method': 'GET', 'status': 'implemented', 'domain': 'inspection_history', 'description': '历史病害检索和原表追溯'},
    {'path': '/api/inspections/history/catalog', 'method': 'GET', 'status': 'implemented', 'domain': 'inspection_history', 'description': '历史巡查来源、日期范围和对象数量'},
    {'path': '/api/inspections/history/source', 'method': 'GET', 'status': 'implemented', 'domain': 'inspection_history', 'description': '下载与导入版本校验一致的原巡查表'},
]


def _raw(value: Any) -> Any:
    return value.isoformat() if isinstance(value, (date, datetime)) else value


def build_inspection_history(paths: DemoPaths = PATHS) -> dict[str, Any]:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    source = paths.root / SOURCE
    raw = source.read_bytes()
    digest = sha256(raw).hexdigest()
    batch_id = sha256(f'{VERSION}:{digest}'.encode()).hexdigest()[:24]
    book = load_workbook(BytesIO(raw), read_only=True, data_only=False)
    rows = []
    skipped_sheets = []
    try:
        for sheet in book:
            iterator = sheet.iter_rows(values_only=True)
            headers = next(iterator, ())
            if sheet.title not in {'边坡', '桥梁'}:
                skipped_sheets.append(sheet.title)
                continue
            keys = [str(h).strip() if h is not None else f'未命名列{get_column_letter(i + 1)}' for i, h in enumerate(headers)]
            if not {'时间', '起止桩号', '方向及位置', '存在问题'} <= set(keys):
                raise ValueError(f'{sheet.title}缺少巡查必要字段')
            for index, values in enumerate(iterator, 2):
                if not any(v is not None for v in values):
                    continue
                original = {k: _raw(v) for k, v in zip(keys, values)}
                day = None
                flags = []
                try:
                    day = datetime.fromisoformat(str(original['时间']).strip()).date().isoformat()
                except (ValueError, TypeError):
                    flags.append('invalid_date')
                if not str(original.get('存在问题') or '').strip():
                    flags.append('missing_problem')
                if any(isinstance(v, str) and v.startswith('=') for v in original.values()):
                    flags.append('formula_preserved')
                rows.append({
                    'id': 'ih_' + sha256(f'{batch_id}:{sheet.title}:{index}'.encode()).hexdigest()[:24],
                    'asset_type': sheet.title, 'observed_date': day, 'raw_time': original['时间'],
                    'station_text': original.get('起止桩号'), 'direction': original.get('方向及位置'),
                    'problem': original.get('存在问题'), 'bridge_name': original.get('桥名'),
                    'raw_fields': original, 'route_id': None, 'slope_id': None,
                    'quality_flags': flags, 'risk_eligible': False,
                    'source': {'file': SOURCE, 'sha256': digest, 'sheet': sheet.title, 'row': index,
                               'cells': {k: f'{get_column_letter(i + 1)}{index}' for i, k in enumerate(keys)}},
                })
    finally:
        book.close()
    if not rows:
        raise ValueError('巡查表没有可归档记录')
    days = sorted(r['observed_date'] for r in rows if r['observed_date'])
    folder = paths.inspection_history_dir / 'snapshots' / batch_id
    events = folder / 'events.json'
    write_json(events, {'rows': rows})
    catalog = {
        'status': 'imported_historical', 'schema_version': VERSION, 'batch_id': batch_id,
        'generated_at': datetime.now(timezone.utc).isoformat(), 'boundary': BOUNDARY,
        'source_file': SOURCE, 'source_sha256': digest, 'count': len(rows),
        'counts_by_type': dict(Counter(r['asset_type'] for r in rows)),
        'start': days[0] if days else None, 'end': days[-1] if days else None,
        'invalid_date_rows': sum(r['observed_date'] is None for r in rows),
        'skipped_sheets': skipped_sheets, 'route_mapping_status': 'unknown', 'risk_eligible': False,
        'events_file': events.relative_to(paths.inspection_history_dir).as_posix(),
        'events_sha256': sha256(events.read_bytes()).hexdigest(),
    }
    write_json(folder / 'catalog.json', catalog)
    write_json(paths.inspection_history_dir / 'catalog.json', catalog)
    return catalog


def history_catalog(paths: DemoPaths = PATHS) -> dict[str, Any]:
    catalog = read_json(paths.inspection_history_dir / 'catalog.json', None)
    if not catalog:
        return {'status': 'not_imported', 'count': 0, 'counts_by_type': {}, 'boundary': BOUNDARY, 'risk_eligible': False}
    file = (paths.inspection_history_dir / catalog['events_file']).resolve()
    file.relative_to(paths.inspection_history_dir.resolve())
    verify_file_sha256(file, catalog.get('events_sha256'))
    return catalog


def public_history_catalog(paths: DemoPaths = PATHS) -> dict[str, Any]:
    return {k: v for k, v in history_catalog(paths).items() if k not in {'events_file', 'events_sha256'}}


def history_payload(query: dict[str, list[str]], paths: DemoPaths = PATHS) -> dict[str, Any]:
    catalog = history_catalog(paths)
    if catalog['status'] == 'not_imported':
        return {**catalog, 'rows': [], 'total': 0}
    asset_type = query.get('asset_type', ['边坡'])[0]
    if asset_type not in {'边坡', '桥梁', 'all'}:
        raise ValueError('对象类型须为边坡、桥梁或all')
    start, end = query.get('start', [''])[0], query.get('end', [''])[0]
    for value in [start, end]:
        if value and date.fromisoformat(value).isoformat() != value:
            raise ValueError('日期须为YYYY-MM-DD')
    if start and end and start > end:
        raise ValueError('开始日期不能晚于结束日期')
    include_undated = query.get('include_undated', ['false'])[0]
    if include_undated not in {'true', 'false'}:
        raise ValueError('include_undated须为true或false')
    limit, offset = int(query.get('limit', ['20'])[0]), int(query.get('offset', ['0'])[0])
    if not 1 <= limit <= 100 or offset < 0:
        raise ValueError('limit须为1—100，offset须为非负整数')
    keyword = query.get('q', [''])[0].strip().casefold()
    if len(keyword) > 200:
        raise ValueError('检索词过长')
    file = paths.inspection_history_dir / catalog['events_file']
    rows = read_json(file, {})['rows']
    def within_dates(row):
        if row['observed_date'] is None:
            return include_undated == 'true' or not (start or end)
        return (not start or row['observed_date'] >= start) and (not end or row['observed_date'] <= end)
    filtered = [r for r in rows if (asset_type == 'all' or r['asset_type'] == asset_type)
                and within_dates(r)
                and (not keyword or keyword in ' '.join(str(r.get(k) or '') for k in ['problem', 'station_text', 'direction', 'bridge_name']).casefold())]
    filtered.sort(key=lambda r: (r['observed_date'] or '', r['source']['sheet'], r['source']['row']))
    return {'status': catalog['status'], 'batch_id': catalog['batch_id'], 'rows': filtered[offset:offset + limit],
            'total': len(filtered), 'count': min(limit, max(0, len(filtered) - offset)),
            'offset': offset, 'limit': limit, 'boundary': BOUNDARY, 'risk_eligible': False}


def history_source(paths: DemoPaths = PATHS) -> tuple[Path, bytes]:
    catalog = history_catalog(paths)
    if catalog['status'] == 'not_imported':
        raise FileNotFoundError('历史巡查尚未导入')
    source = (paths.root / catalog['source_file']).resolve()
    if source != (paths.root / SOURCE).resolve():
        raise ValueError('巡查来源路径不合法')
    data = source.read_bytes()
    if sha256(data).hexdigest() != catalog['source_sha256']:
        raise OSError('原巡查表已变化，请重新导入后下载')
    return source, data


def history_interfaces(payload: dict[str, Any], paths: DemoPaths = PATHS) -> dict[str, Any]:
    try:
        catalog = history_catalog(paths)
        status = 'implemented_historical' if catalog.get('count') else 'implemented_not_imported'
        description = '历史病害检索和原表追溯；近期巡检与正式边坡关联仍待补'
    except (OSError, ValueError, KeyError, TypeError) as exc:
        status, description = 'data_unavailable', '历史巡查暂不可读：' + str(exc)
    existing = {r['path'] for r in payload.get('interfaces', [])}
    registered = [*payload.get('interfaces', []), *(r for r in HISTORY_INTERFACES if r['path'] not in existing)]
    rows = [{**r, 'status': status, 'description': description} if r.get('path') == '/api/inspections/history' else dict(r)
            for r in registered]
    return {**payload, 'interfaces': rows, 'status_counts': dict(Counter(r['status'] for r in rows))}
